# -*- coding: utf-8 -*-
"""差异报告输出：生成 Excel 报告（多 sheet + 汇总 + 明细）。"""

from __future__ import annotations

import os
from collections import Counter
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .rules import CheckResult, Diff, NamePair

_HDR_FILL = PatternFill("solid", fgColor="D9E2F3")
_HARD_FILL = PatternFill("solid", fgColor="FDE9E9")   # 红底
_SUS_FILL = PatternFill("solid", fgColor="FFF6E0")    # 黄底
_THIN = Side(style="thin", color="BFBFBF")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _style_header(ws, row: int, ncols: int):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(bold=True)
        cell.fill = _HDR_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _BORDER


def _write_rows(ws, rows: list[list], start_row: int, ncols: int, level_col: int | None = None):
    for i, row in enumerate(rows, start=start_row):
        for c, v in enumerate(row, start=1):
            cell = ws.cell(row=i, column=c, value=v)
            cell.border = _BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if level_col and c == level_col:
                if v == "硬错误":
                    cell.fill = _HARD_FILL
                elif v == "可疑项":
                    cell.fill = _SUS_FILL


def _auto_width(ws, ncols: int, min_w: int = 10, max_w: int = 40):
    for c in range(1, ncols + 1):
        letter = get_column_letter(c)
        longest = 0
        for cell in ws[letter]:
            if cell.value is not None:
                longest = max(longest, len(str(cell.value)))
        ws.column_dimensions[letter].width = max(min_w, min(longest + 2, max_w))


def build_report(result: CheckResult, period: tuple[date, date],
                 out_path: str, name_pairs: list[NamePair] | None = None,
                 doc_records: list | None = None) -> str:
    os.makedirs(os.path.dirname(out_path) or ".", exist_ok=True)
    wb = Workbook()

    # ---- Sheet 1: 汇总 ----
    ws = wb.active
    ws.title = "差异汇总"
    by_level = result.by_level()
    by_cat = result.by_category()
    ws.append(["考勤对账差异汇总（阶段二：原始打卡 × 手工考勤表 × PDF单据）"])
    ws.append([f"对账周期：{period[0].isoformat()} ~ {period[1].isoformat()}"])
    ws.append([f"差异总数：{len(result.diffs)}（硬错误 {by_level.get('硬错误', 0)}，可疑项 {by_level.get('可疑项', 0)}）"])
    ws.append([])
    ws.append(["差异类别", "数量"])
    for cat, n in sorted(by_cat.items(), key=lambda kv: -kv[1]):
        ws.append([cat, n])
    ws.append(["合计", len(result.diffs)])
    ws.merge_cells("A1:B1")
    ws.merge_cells("A2:B2")
    ws.merge_cells("A3:B3")
    _style_header(ws, 5, 2)
    _auto_width(ws, 2)

    # ---- Sheet 2: 到场核对明细 ----
    ws2 = wb.create_sheet("到场核对明细")
    heads = ["类别", "姓名", "日期", "级别", "原始记录", "考勤表记录", "差异说明"]
    ws2.append(heads)
    _style_header(ws2, 1, len(heads))
    rows = [d.to_row() for d in result.diffs
            if d.category in ("到场核对", "加班核对", "缺卡提醒", "假单核对(待阶段二)")]
    _write_rows(ws2, rows, 2, len(heads), level_col=4)
    ws2.auto_filter.ref = f"A1:G{max(1, len(rows) + 1)}"
    _auto_width(ws2, len(heads))

    # ---- Sheet 3: 累计重算明细 ----
    ws3 = wb.create_sheet("累计重算明细")
    ws3.append(heads)
    _style_header(ws3, 1, len(heads))
    rows3 = [d.to_row() for d in result.diffs if d.category == "累计重算"]
    _write_rows(ws3, rows3, 2, len(heads), level_col=4)
    ws3.auto_filter.ref = f"A1:G{max(1, len(rows3) + 1)}"
    _auto_width(ws3, len(heads))

    # ---- Sheet 4: 姓名配对 ----
    ws5 = wb.create_sheet("姓名配对")
    heads5 = ["姓名(简体)", "配对状态", "原始记录", "考勤表", "工号", "部门", "建议配对(相似度)"]
    ws5.append(heads5)
    _style_header(ws5, 1, len(heads5))
    rows5 = []
    if name_pairs:
        for p in name_pairs:
            rows5.append(p.to_row())
    _write_rows(ws5, rows5, 2, len(heads5))
    ws5.auto_filter.ref = f"A1:G{max(1, len(rows5) + 1)}"
    _auto_width(ws5, len(heads5))
    # 配对状态着色
    for i, p in enumerate(name_pairs or [], start=2):
        if p.status == "仅考勤表":
            ws5.cell(row=i, column=2).fill = PatternFill("solid", fgColor="FDE9E9")
        elif p.status == "仅原始记录":
            ws5.cell(row=i, column=2).fill = PatternFill("solid", fgColor="FFF6E0")

    # ---- Sheet 5: 单据核对明细 ----
    ws_doc = wb.create_sheet("单据核对明细")
    heads_doc = ["类别", "姓名", "日期", "级别", "原始记录", "考勤表记录", "差异说明"]
    ws_doc.append(heads_doc)
    _style_header(ws_doc, 1, len(heads_doc))
    rows_doc = [d.to_row() for d in result.diffs if d.category == "单据核对"]
    _write_rows(ws_doc, rows_doc, 2, len(heads_doc), level_col=4)
    ws_doc.auto_filter.ref = f"A1:G{max(1, len(rows_doc) + 1)}"
    _auto_width(ws_doc, len(heads_doc))

    # ---- Sheet 6: 单据清单 ----
    ws_list = wb.create_sheet("单据清单")
    heads_list = ["页码", "单据类型", "员工", "部门", "假期/加班类型", "开始日期", "结束日期",
                  "天数", "加班日期", "加班时间", "加班时数", "置信度", "备注"]
    ws_list.append(heads_list)
    _style_header(ws_list, 1, len(heads_list))
    rows_list = []
    if doc_records:
        for r in doc_records:
            rows_list.append([
                r.page_num, r.doc_type_label, r.employee_name or "", r.department or "",
                r.leave_type_label or "",
                r.start_date.isoformat() if r.start_date else "",
                r.end_date.isoformat() if r.end_date else "",
                r.days if r.days is not None else "",
                r.overtime_date.isoformat() if r.overtime_date else "",
                f"{r.overtime_start or ''}~{r.overtime_end or ''}" if r.overtime_start else "",
                r.overtime_hours if r.overtime_hours is not None else "",
                r.confidence, r.notes,
            ])
    _write_rows(ws_list, rows_list, 2, len(heads_list))
    ws_list.auto_filter.ref = f"A1:M{max(1, len(rows_list) + 1)}"
    _auto_width(ws_list, len(heads_list))
    # 置信度着色
    for i, r in enumerate(doc_records or [], start=2):
        if r.confidence == "low":
            ws_list.cell(row=i, column=12).fill = PatternFill("solid", fgColor="FDE9E9")
        elif r.confidence == "medium":
            ws_list.cell(row=i, column=12).fill = PatternFill("solid", fgColor="FFF6E0")

    # ---- Sheet 7: 全部差异 ----
    ws4 = wb.create_sheet("全部差异")
    ws4.append(heads)
    _style_header(ws4, 1, len(heads))
    rows4 = [d.to_row() for d in result.diffs]
    _write_rows(ws4, rows4, 2, len(heads), level_col=4)
    ws4.auto_filter.ref = f"A1:G{max(1, len(rows4) + 1)}"
    _auto_width(ws4, len(heads))

    return save_with_fallback(wb, out_path)


def save_with_fallback(wb: Workbook, out_path: str) -> str:
    """保存工作簿；文件被占用（Excel 打开中）时自动加时间戳另存。"""
    from datetime import datetime

    try:
        wb.save(out_path)
        return out_path
    except PermissionError:
        base, ext = os.path.splitext(out_path)
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        alt = f"{base}_{stamp}{ext}"
        wb.save(alt)
        return alt
